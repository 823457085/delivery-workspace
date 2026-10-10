#!/usr/bin/env python3
"""
Export selected locale keys from zh-CN.json to an Excel workbook.

Default columns:
  多端类型 | 位置编码 | 页面ID | 原文 | cn_text | en_text | fr_text | de_text | es_text | ja_text | pt_text
"""

from __future__ import annotations

import argparse
import json
import re
import sys
from copy import copy
from pathlib import Path

try:
    import openpyxl
    from openpyxl.styles import Alignment, Font, PatternFill
    from openpyxl.utils import get_column_letter
except ImportError:
    print("请安装: pip3 install openpyxl", file=sys.stderr)
    sys.exit(1)

SKILL_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_TEMPLATE_PATH = SKILL_ROOT / "templates" / "多语言模板-后台导入.xlsx"
SHEET_NAME = "硬编码扫描结果"
TEMPLATE_HEADERS = ["多端类型", "位置编码", "页面ID", "原文", "cn_text", "en_text", "fr_text", "de_text", "es_text", "ja_text", "pt_text"]
LEGACY_HEADERS = ["多端类型", "位置编码", "页面ID", "原文"]
HTML_TAG_RE = re.compile(r"</?[A-Za-z][^>]*>")


def extract_terminal_name(project_root: Path) -> str:
    locale_file = project_root / "src" / "locale" / "index.js"
    if not locale_file.exists():
        return "APEX"
    text = locale_file.read_text(encoding="utf-8")
    match = re.search(r"TERMINALNAME\s*=\s*[\"']([^\"']+)[\"']", text)
    return match.group(1) if match else "APEX"


def page_id_from_code(terminal: str, code: str) -> str:
    if not code.upper().startswith(terminal.upper()):
        return ""
    rest = code[len(terminal) :]
    return (rest[:4] if len(rest) >= 4 else rest).upper()


def parse_keys(keys: str | None, keys_file: Path | None) -> list[str]:
    raw_parts: list[str] = []
    if keys:
        raw_parts.append(keys)
    if keys_file:
        if not keys_file.exists():
            print(f"未找到 keys 文件: {keys_file}", file=sys.stderr)
            sys.exit(1)
        raw_parts.append(keys_file.read_text(encoding="utf-8"))
    parsed: list[str] = []
    for raw in raw_parts:
        parsed.extend(k.strip().upper() for k in re.split(r"[\s,]+", raw) if k.strip())
    return list(dict.fromkeys(parsed))


def resolve_project_path(project_root: Path, value: Path | None) -> Path | None:
    if value is None:
        return None
    return value if value.is_absolute() else project_root / value


def read_json_dict(path: Path, label: str) -> dict[str, str]:
    if not path.exists():
        print(f"未找到 {label}: {path}", file=sys.stderr)
        sys.exit(1)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        print(f"{label} 不是合法 JSON: {path} ({exc})", file=sys.stderr)
        sys.exit(1)
    if not isinstance(data, dict):
        print(f"{label} 必须是 JSON 对象: {path}", file=sys.stderr)
        sys.exit(1)
    return {str(key).upper(): str(value) for key, value in data.items() if value is not None}


def read_translation_map(path: Path) -> dict[str, str]:
    if not path.exists():
        print(f"未找到英文翻译映射: {path}", file=sys.stderr)
        sys.exit(1)
    try:
        data = json.loads(path.read_text(encoding="utf-8"))
    except json.JSONDecodeError as exc:
        print(f"英文翻译映射不是合法 JSON: {path} ({exc})", file=sys.stderr)
        sys.exit(1)
    if not isinstance(data, dict):
        print(f"英文翻译映射必须是 JSON 对象: {path}", file=sys.stderr)
        sys.exit(1)
    result: dict[str, str] = {}
    for key, value in data.items():
        if value is not None:
            result[str(key)] = str(value)
            result[str(key).upper()] = str(value)
    return result


def find_default_english_locale(project_root: Path) -> Path | None:
    for relative in ("src/locale/en.json", "src/locale/en-US.json"):
        path = project_root / relative
        if path.exists():
            return path
    return None


def load_english_sources(
    project_root: Path,
    english_locale: Path | None,
    english_map: Path | None,
) -> tuple[dict[str, str], dict[str, str]]:
    english_locale_path = resolve_project_path(project_root, english_locale) if english_locale else find_default_english_locale(project_root)
    english_messages = read_json_dict(english_locale_path, "英文 locale") if english_locale_path else {}
    english_overrides = read_translation_map(resolve_project_path(project_root, english_map)) if english_map else {}
    return english_messages, english_overrides


def resolve_english_text(
    key: str,
    chinese_text: str,
    english_messages: dict[str, str],
    english_overrides: dict[str, str],
) -> str:
    for source in (english_overrides, english_messages):
        value = source.get(key) or source.get(key.upper()) or source.get(chinese_text)
        if value:
            return value
    return ""


def terminal_display_name(terminal: str) -> str:
    return terminal[:1].upper() + terminal[1:].lower() if terminal else terminal


def capture_cell_style(cell) -> dict[str, object]:
    return {
        "font": copy(cell.font),
        "fill": copy(cell.fill),
        "border": copy(cell.border),
        "alignment": copy(cell.alignment),
        "number_format": cell.number_format,
        "protection": copy(cell.protection),
    }


def apply_cell_style(cell, style: dict[str, object]) -> None:
    cell.font = copy(style["font"])
    cell.fill = copy(style["fill"])
    cell.border = copy(style["border"])
    cell.alignment = copy(style["alignment"])
    cell.number_format = style["number_format"]
    cell.protection = copy(style["protection"])


def load_template(template_path: Path) -> tuple[openpyxl.Workbook, object, list[dict[str, object]], float | None]:
    if not template_path.exists():
        print(f"未找到导入模板: {template_path}", file=sys.stderr)
        sys.exit(1)
    wb = openpyxl.load_workbook(template_path)
    ws = wb[SHEET_NAME] if SHEET_NAME in wb.sheetnames else wb.active
    ws.title = SHEET_NAME
    headers = [ws.cell(row=1, column=col).value for col in range(1, len(TEMPLATE_HEADERS) + 1)]
    if headers != TEMPLATE_HEADERS:
        print(f"模板表头不符合预期: {headers}", file=sys.stderr)
        print(f"预期表头: {TEMPLATE_HEADERS}", file=sys.stderr)
        sys.exit(1)
    data_styles = [capture_cell_style(ws.cell(row=2, column=col)) for col in range(1, len(TEMPLATE_HEADERS) + 1)]
    data_row_height = ws.row_dimensions[2].height
    if ws.max_row > 1:
        ws.delete_rows(2, ws.max_row - 1)
    return wb, ws, data_styles, data_row_height


def build_legacy_workbook(terminal: str, keys: list[str], messages: dict[str, str]) -> openpyxl.Workbook:
    wb = openpyxl.Workbook()
    ws = wb.active
    ws.title = SHEET_NAME

    header_fill = PatternFill(start_color="4472C4", end_color="4472C4", fill_type="solid")
    header_font = Font(color="FFFFFF", bold=True, size=11)
    center = Alignment(horizontal="center", vertical="center")

    for col, header in enumerate(LEGACY_HEADERS, 1):
        cell = ws.cell(row=1, column=col, value=header)
        cell.fill = header_fill
        cell.font = header_font
        cell.alignment = center

    for col, width in enumerate([12, 18, 10, 48], 1):
        ws.column_dimensions[get_column_letter(col)].width = width

    for row, key in enumerate(keys, 2):
        ws.cell(row=row, column=1, value=terminal)
        ws.cell(row=row, column=2, value=key)
        ws.cell(row=row, column=3, value=page_id_from_code(terminal, key))
        ws.cell(row=row, column=4, value=messages.get(key, ""))

    return wb


def build_template_workbook(
    terminal: str,
    terminal_display: str,
    keys: list[str],
    chinese_messages: dict[str, str],
    english_messages: dict[str, str],
    english_overrides: dict[str, str],
    template_path: Path,
    allow_missing_english: bool,
) -> openpyxl.Workbook:
    wb, ws, data_styles, data_row_height = load_template(template_path)
    missing_english: list[tuple[str, str]] = []
    invalid_original: list[tuple[str, str]] = []

    for row, key in enumerate(keys, 2):
        chinese_text = chinese_messages.get(key, "")
        english_text = resolve_english_text(key, chinese_text, english_messages, english_overrides)
        if chinese_text and not english_text:
            missing_english.append((key, chinese_text))
        if english_text and HTML_TAG_RE.search(english_text):
            invalid_original.append((key, english_text))

        values = [
            terminal_display,
            key,
            page_id_from_code(terminal, key),
            english_text,
            chinese_text,
            english_text,
            "",
            "",
            "",
            "",
            "",
        ]
        for col, value in enumerate(values, 1):
            cell = ws.cell(row=row, column=col, value=value)
            apply_cell_style(cell, data_styles[col - 1])
        if data_row_height is not None:
            ws.row_dimensions[row].height = data_row_height

    if missing_english and not allow_missing_english:
        print("以下 Key 缺少英文翻译，未生成 xlsx。请提供 --english-locale 或 --english-map：", file=sys.stderr)
        for key, chinese_text in missing_english[:50]:
            print(f"  {key}: {chinese_text}", file=sys.stderr)
        if len(missing_english) > 50:
            print(f"  ... 还有 {len(missing_english) - 50} 条", file=sys.stderr)
        sys.exit(1)

    if invalid_original:
        print("以下 Key 的英文原文包含 HTML-like 标签，未生成 xlsx：", file=sys.stderr)
        for key, original_text in invalid_original[:50]:
            print(f"  {key}: {original_text}", file=sys.stderr)
        sys.exit(1)

    return wb


def main() -> None:
    parser = argparse.ArgumentParser(description="导出指定 i18n Key 为多语言 Excel")
    parser.add_argument("--project-root", default=".", help="项目根目录，默认当前目录")
    parser.add_argument("--keys", help="逗号或空白分隔的 Key 列表")
    parser.add_argument("--keys-file", type=Path, help="包含 Key 的文件，支持逗号/空白/换行分隔")
    parser.add_argument("-o", "--output", default="tempdocs/locale-keys-export.xlsx", help="输出 xlsx 路径")
    parser.add_argument("--locale", default="src/locale/zh-CN.json", help="中文 JSON 路径（相对项目根）")
    parser.add_argument("--terminal", help="端类型；默认从 src/locale/index.js 读取，读不到则为 APEX")
    parser.add_argument("--terminal-display", help="写入“多端类型”的展示值；默认如 APEX -> Apex")
    parser.add_argument("--template", type=Path, default=DEFAULT_TEMPLATE_PATH, help="后台批量导入模板 xlsx")
    parser.add_argument("--english-locale", type=Path, help="英文 locale JSON，支持相对项目根路径；未指定时尝试 src/locale/en.json 或 en-US.json")
    parser.add_argument("--english-map", type=Path, help="英文翻译映射 JSON，支持 key->英文 或 中文->英文")
    parser.add_argument("--allow-missing-english", action="store_true", help="允许缺少英文翻译时导出空值")
    parser.add_argument("--legacy-format", action="store_true", help="使用旧 4 列格式：多端类型/位置编码/页面ID/原文")
    args = parser.parse_args()

    project_root = Path(args.project_root).resolve()
    locale_path = project_root / args.locale
    if not locale_path.exists():
        print(f"未找到 locale 文件: {locale_path}", file=sys.stderr)
        sys.exit(1)

    keys = parse_keys(args.keys, args.keys_file)
    if not keys:
        print("请通过 --keys 或 --keys-file 提供至少一个 Key", file=sys.stderr)
        sys.exit(1)

    terminal = args.terminal or extract_terminal_name(project_root)
    messages = read_json_dict(locale_path, "中文 locale")
    if args.legacy_format:
        wb = build_legacy_workbook(terminal, keys, messages)
    else:
        english_messages, english_overrides = load_english_sources(project_root, args.english_locale, args.english_map)
        template_path = args.template if args.template.is_absolute() else SKILL_ROOT / args.template
        wb = build_template_workbook(
            terminal=terminal,
            terminal_display=args.terminal_display or terminal_display_name(terminal),
            keys=keys,
            chinese_messages=messages,
            english_messages=english_messages,
            english_overrides=english_overrides,
            template_path=template_path,
            allow_missing_english=args.allow_missing_english,
        )

    output = Path(args.output)
    if not output.is_absolute():
        output = project_root / output
    output.parent.mkdir(parents=True, exist_ok=True)
    wb.save(str(output))
    print(f"已生成: {output}（共 {len(keys)} 条）")


if __name__ == "__main__":
    main()
